"""Print what is really in a deck: per slide its kind tag, layout, every shape (id, name, kind, position and size in
inches, fill, text with font size / bold / color / typeface), the click steps (which shape ids each click shows) and the
start of the notes. Use it to look at the deck you copy from before you edit it.
Usage: python3 scripts/slides/inspect_deck.py <deck.pptx> [slide numbers ...] [--short]"""
import sys, re, zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Emu
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
args = [a for a in sys.argv[1:] if not a.startswith('--')]
short = '--short' in sys.argv
prs = Presentation(args[0]); want = {int(a) for a in args[1:]}
inch = lambda v: f'{Emu(v).inches:.2f}'
print(f'{args[0]}: {len(prs.slides)} slides, {inch(prs.slide_width)} x {inch(prs.slide_height)} in')
def fill_of(el):
    c = el.find('p:spPr/a:solidFill/a:srgbClr', NS)
    return c.get('val') if c is not None else ('none' if el.find('p:spPr/a:noFill', NS) is not None else '-')
def runs(el):
    out = []
    for r in el.iterfind('.//a:r', NS):
        rp = r.find('a:rPr', NS); t = r.findtext('a:t', '', NS)
        if not t.strip(): continue
        sz = int(rp.get('sz', '0')) / 100 if rp is not None else 0
        col = rp.find('a:solidFill/a:srgbClr', NS) if rp is not None else None
        face = rp.find('a:latin', NS) if rp is not None else None
        out.append(f"[{sz:g}pt{' b' if rp is not None and rp.get('b') == '1' else ''}"
                   f"{' #' + col.get('val') if col is not None else ''}{' ' + face.get('typeface') if face is not None else ''}] {t[:60]}")
    return out
for i, s in enumerate(prs.slides, 1):
    if want and i not in want: continue
    tag = s._element.find('p:cSld', NS).get('name') or ''
    print(f'\n=== slide {i}  tag={tag!r}  layout={s.slide_layout.name!r}')
    for el in s.shapes._spTree:
        k = etree.QName(el).localname
        if k not in ('sp', 'pic', 'grpSp', 'cxnSp'): continue
        c = el.find('.//p:cNvPr', NS); off = el.find('.//a:off', NS); ext = el.find('.//a:ext', NS)
        geo = el.find('.//a:prstGeom', NS)
        pos = f'({inch(int(off.get("x")))},{inch(int(off.get("y")))}) {inch(int(ext.get("cx")))}x{inch(int(ext.get("cy")))}' if off is not None else ''
        line = f'  id={c.get("id"):>3} {k:4} {c.get("name")!r:28} {pos:28} {geo.get("prst") if geo is not None else "":10} fill={fill_of(el) if k == "sp" else "-"}'
        if k == 'pic': line += f'  alt={(c.get("descr") or "")[:50]!r}'
        print(line)
        if not short:
            for r in runs(el): print('        ' + r)
    steps = []
    for ctn in s._element.iterfind('.//p:timing//p:cTn[@nodeType]', NS):
        if ctn.get('nodeType') in ('clickEffect', 'withEffect', 'afterEffect'):
            ids = sorted({int(t.get('spid')) for t in ctn.iterfind('.//p:spTgt', NS)})
            steps.append(('+' if ctn.get('nodeType') != 'clickEffect' else '') + ','.join(map(str, ids)))
    print('  clicks:', ' | '.join(steps) or 'none')
    if s.has_notes_slide:
        n = s.notes_slide.notes_text_frame.text
        print('  notes:', (n[:300] + '…') if len(n) > 300 else n.replace('\n', ' / '))
